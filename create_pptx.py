"""
生成 PowerPoint 文件
"""
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

def create_pptx():
    # 读取 pages.json
    pages_file = Path("d:/ragku/frontend/public/assets/images/posters/pages.json")
    with open(pages_file, 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    pages = data['pages']
    
    # 创建演示文稿 (16:9)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # 空白布局
    blank_layout = prs.slide_layouts[6]
    
    # 截图目录
    screenshot_dir = Path("d:/ragku/frontend/public/assets/images/posters/pages")
    
    for page in pages:
        page_num = page['number']
        screenshot_file = screenshot_dir / f"page-{page_num:02d}.png"
        
        if screenshot_file.exists():
            # 添加幻灯片
            slide = prs.slides.add_slide(blank_layout)
            
            # 添加图片 (全屏)
            left = Inches(0)
            top = Inches(0)
            width = prs.slide_width
            height = prs.slide_height
            
            slide.shapes.add_picture(
                str(screenshot_file),
                left, top, width, height
            )
            
            print(f"已添加第 {page_num} 页: {page['title']}")
        else:
            print(f"警告: 第 {page_num} 页的截图不存在")
    
    # 保存 PPT
    output_file = Path("d:/ragku/RAG项目展示.pptx")
    prs.save(str(output_file))
    print(f"\nPPT 已保存到: {output_file}")
    return output_file

if __name__ == "__main__":
    try:
        from pptx import Presentation
    except ImportError:
        print("正在安装 python-pptx...")
        import subprocess
        import sys
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx", "-q"])
        from pptx import Presentation
    
    create_pptx()
